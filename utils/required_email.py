"""Mandatory evidence for email deliverables, separate from partial-credit score."""
from io import BytesIO
from zipfile import ZipFile


def readable_attachment(filename, content_type, content):
    """Validate the attached bytes; never fall back to a workspace/content_path file.

    The benchmark PG email backend uses application/octet-stream for all uploads.
    Accept that MIME as well as the correct OOXML MIME, but parse the actual file.
    """
    formats = {
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }
    suffix = next((ext for ext in formats if filename.endswith(ext)), None)
    if suffix is None or content_type not in (formats[suffix], "application/octet-stream"):
        return False
    if not isinstance(content, (bytes, bytearray, memoryview)) or not content:
        return False
    raw = bytes(content)
    try:
        with ZipFile(BytesIO(raw)) as package:
            if package.testzip() is not None:
                return False
        if suffix == ".docx":
            from docx import Document
            Document(BytesIO(raw))
        else:
            from openpyxl import load_workbook
            book = load_workbook(BytesIO(raw), data_only=True)
            book.close()
    except ImportError:
        raise
    except Exception:
        return False
    return True


def required_sent_email(conn, recipient, filename):
    """Return sent subject/body rows and the two mandatory evidence flags.

    email.sent_log is the benchmark's send receipt (not proof of external SMTP
    delivery). An arbitrary email.messages row or local file is insufficient.
    Recipient matching is exact and case insensitive on the JSONB recipient list.
    An attachment must belong to one of those same sent messages.
    """
    cur = conn.cursor()
    try:
        cur.execute(
            """SELECT DISTINCT m.id, m.subject, m.body_text
                 FROM email.messages m
                 JOIN email.sent_log s ON s.message_id = m.id
                WHERE EXISTS (
                    SELECT 1 FROM jsonb_array_elements_text(m.to_addr) AS r(address)
                     WHERE lower(r.address) = lower(%s)
                )""",
            (recipient,),
        )
        rows = cur.fetchall()
        attached = False
        if rows:
            cur.execute(
                """SELECT filename, content_type, content FROM email.attachments
                    WHERE message_id = ANY(%s) AND filename = %s""",
                ([row[0] for row in rows], filename),
            )
            attached = any(
                name == filename and readable_attachment(name, mime, content)
                for name, mime, content in cur.fetchall()
            )
        return ([(subject, body) for _id, subject, body in rows],
                {"email_sent": bool(rows), "email_attachment": attached})
    finally:
        cur.close()
