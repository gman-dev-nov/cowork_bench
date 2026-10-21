"""Fail-closed gates for mandatory task artifacts.

Some benchmark deliverables are mandatory Word documents: a task cannot pass
if the required .docx is missing or unreadable, even when enough soft checks
succeed to push accuracy over the pass threshold. Graders use
``require_readable_docx`` as an early gate, immediately after argument
parsing and before any DB reads/writes or component checks.
"""
import sys

from docx import Document


def require_readable_docx(path):
    """Fail closed unless ``path`` is a readable DOCX (real ZIP/OOXML).

    Parses the mandatory file with the same python-docx ``Document`` parser
    the graders use for semantic checks. On any problem (missing file,
    unreadable bytes, plain text renamed to .docx, malformed or non-DOCX
    zip) prints a concise CRITICAL diagnostic and raises SystemExit(1).
    On success returns the parsed ``Document``. Never writes or rewrites
    the input file.
    """
    try:
        return Document(path)
    except Exception as exc:  # missing, not a zip, corrupt, wrong format...
        print(f"CRITICAL: required Word document is missing or unreadable: {path}")
        print(f"CRITICAL: {type(exc).__name__}: {exc}")
        print("FAIL (required document gate)")
        sys.exit(1)


def require_readable_xlsx(path):
    """Mandatory local spreadsheets must be real readable workbooks."""
    from openpyxl import load_workbook
    try:
        book = load_workbook(path, read_only=True, data_only=True)
        book.close()
    except Exception as exc:
        print(f"CRITICAL: required spreadsheet is missing or unreadable: {path}")
        print(f"CRITICAL: {type(exc).__name__}: {exc}")
        raise SystemExit(1)
