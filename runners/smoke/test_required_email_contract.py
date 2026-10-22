"""Actual three graders/main, real Office fixtures, isolated service DB state."""
from pathlib import Path
from io import BytesIO, StringIO
import contextlib
import datetime as dt
import importlib.util
import json
import sys
import tempfile
import unittest
from unittest.mock import patch
from docx import Document
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
SPECS = [
    ('rzd-canvas-fieldtrip-novgorod-gcal-word-email', 'students@university.ru',
     'Field_Trip_Notice.docx', 'Учебная экскурсия', '818А 819А 5000'),
    ('rzd-kulinar-team-trip-spb-catering-excel-gcal', 'events@company.ru',
     'Team_Trip_Plan.xlsx', 'План поездки Санкт-Петербург', '752А Сапсан план Санкт-Петербург'),
    ('rzd-msk-spb-team-trip-gcal-word', 'travel@consulting.ru',
     'Travel_Plan.docx', 'Деловая поездка', '752А 759А 3 человека 33000'),
]


def workspace(folder):
    for filename, headings, text in [
        ('Field_Trip_Notice.docx', ['Описание поездки', 'Расписание', 'Требования к студентам', 'Расходы'],
         'Москва Великий Новгород 818А 819А 5000'),
        ('Travel_Plan.docx', ['Поездка туда', 'Поездка обратно', 'Сводка'],
         'Москва Санкт-Петербург 752А 759А 06:50 10:50 19:25 23:25 5500 16500 33000'),
    ]:
        doc = Document()
        for heading in headings:
            doc.add_heading(heading, 1)
        doc.add_paragraph(text)
        for _ in range(2):
            doc.add_table(rows=1, cols=1).cell(0, 0).text = text
        doc.save(folder / filename)
    (folder / 'trip_notes.md').write_text('752А 759А 2026-03-10 3 человека 33000')
    book = Workbook()
    book.active.title = 'Travel'
    book.active.append(['Train_No', 'Departure'])
    book.active.append(['752А', '06:50'])
    menu = book.create_sheet('Menu')
    menu.append(['Course_Type', 'Dish_Name'])
    timeline = book.create_sheet('Timeline')
    timeline.append(['Time', 'Activity'])
    for i in range(6):
        menu.append(['main', 'synthetic dish'])
        timeline.append([str(i), 'synthetic activity'])
    book.save(folder / 'Team_Trip_Plan.xlsx')
    book.close()


class ServiceDB:
    """Narrow DB seam; real SQL is checked separately on disposable PostgreSQL.

    An unsent messages row remains visible to the old fallback query. Attached
    files are selected by the query's message IDs and filename, never from disk.
    """
    def __init__(self, spec, raw, scenario):
        self.task, self.recipient, self.filename, self.subject, self.body = spec
        self.present = scenario != 'missing'
        self.sent = scenario != 'unsent'
        if scenario == 'partial-body':
            self.body = 'unrelated content'
        self.attachments = []
        if scenario != 'missing-attachment':
            filename = 'Other'+self.filename if scenario == 'wrong-name' else self.filename
            suffix = 'spreadsheetml.sheet' if self.filename.endswith('xlsx') else 'wordprocessingml.document'
            mime = 'application/octet-stream' if scenario == 'generic-mime' else \
                'application/vnd.openxmlformats-officedocument.'+suffix
            if scenario == 'wrong-mime':
                mime = 'text/plain'
            if scenario == 'corrupt':
                raw = b'PK\\x03\\x04truncated'
            if scenario == 'renamed-text':
                raw = b'text is not an Office package'
            self.attachments = [(2 if scenario == 'unrelated-attachment' else 1,
                                 filename, mime, memoryview(raw))]
        self.queries = []
        self.closed = False
    def connect(self, **kwargs):
        return self
    def cursor(self):
        return Cursor(self)
    def close(self):
        self.closed = True


class Cursor:
    def __init__(self, db):
        self.db = db
        self.rows = []
    def execute(self, query, params=None):
        self.db.queries.append((query, params))
        if 'FROM email.messages' in query:
            available = self.db.present and ('email.sent_log' not in query or self.db.sent)
            self.rows = [] if not available else [(self.db.subject, self.db.body)]
            if available and 'm.id,' in query:
                self.rows = [(1, self.db.subject, self.db.body)]
        elif 'FROM email.attachments' in query:
            ids, filename = params
            self.rows = [(name, mime, data) for mid, name, mime, data in self.db.attachments
                         if mid in ids and name == filename]
        elif 'FROM canvas.announcements' in query:
            self.rows = [(1,)]
        elif 'FROM gcal.events' in query:
            dates = [12, 15] if '2026-03-12' in query else [10, 10, 10]
            labels = ['Москва Санкт-Петербург отправление', 'прибытие заселение', 'ужин']
            four = 'description' in query
            three = 'end_datetime' in query and not four
            self.rows = []
            for i, day in enumerate(dates):
                start = dt.datetime(2026, 3, day, 10)
                end = dt.datetime(2026, 3, day, 11)
                self.rows.append((labels[i], 'trip', start, end) if four else
                                 (labels[i], start, end) if three else (labels[i], start))
        else:
            raise AssertionError('Unexpected SQL: '+query)
    def fetchall(self):
        return self.rows
    def fetchone(self):
        return self.rows[0]
    def close(self):
        pass


class RequiredEmailTests(unittest.TestCase):
    def evaluate(self, spec, scenario, low_score=False):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            workspace(folder)
            raw = (folder / spec[2]).read_bytes()
            if scenario == 'wrong-package':
                other = 'Travel_Plan.docx' if spec[2].endswith('xlsx') else 'Team_Trip_Plan.xlsx'
                raw = (folder / other).read_bytes()
            db = ServiceDB(spec, raw, scenario)
            path = ROOT / 'tasks/finalpool' / spec[0] / 'evaluation/main.py'
            module_spec = importlib.util.spec_from_file_location('email_gate_grader', path)
            grader = importlib.util.module_from_spec(module_spec)
            module_spec.loader.exec_module(grader)
            if low_score:
                def failures(*args):
                    for _ in range(15):
                        grader.record('isolated non-email failure', False)
                for name in ['check_word_doc', 'check_trip_notes', 'check_excel',
                             'check_gcal', 'check_canvas_announcement']:
                    if hasattr(grader, name):
                        setattr(grader, name, failures)
            result_file = folder / 'result.json'
            stdout = StringIO()
            with patch.object(grader.psycopg2, 'connect', side_effect=db.connect), \
                 patch.object(sys, 'argv', ['grader', '--agent_workspace', str(folder),
                                          '--res_log_file', str(result_file)]), \
                 contextlib.redirect_stdout(stdout):
                with self.assertRaises(SystemExit) as exited:
                    grader.main()
            result = json.loads(result_file.read_text())
            return exited.exception.code, result, stdout.getvalue()

    def assert_scenario(self, scenario, expected):
        for spec in SPECS:
            with self.subTest(task=spec[0], scenario=scenario):
                code, result, text = self.evaluate(spec, scenario)
                self.assertGreaterEqual(result['accuracy'], 70)
                self.assertEqual(code == 0, expected, text)
                self.assertEqual(result['pass'], expected)
                self.assertEqual(result['mandatory_pass'], expected)
                self.assertEqual(result['partial_score'], result['accuracy'])
                self.assertIn('mandatory_checks', result)

    def test_missing_email_cannot_be_offset_by_other_deliverables(self):
        self.assert_scenario('missing', False)
    def test_unsent_messages_row_is_not_send_evidence(self):
        self.assert_scenario('unsent', False)
    def test_sent_message_without_attachment_fails(self):
        self.assert_scenario('missing-attachment', False)
    def test_wrong_attachment_name_fails(self):
        self.assert_scenario('wrong-name', False)
    def test_wrong_attachment_mime_fails(self):
        self.assert_scenario('wrong-mime', False)
    def test_corrupt_package_fails(self):
        self.assert_scenario('corrupt', False)
    def test_renamed_text_fails(self):
        self.assert_scenario('renamed-text', False)
    def test_wrong_office_package_type_fails(self):
        self.assert_scenario('wrong-package', False)
    def test_attachment_on_another_message_fails(self):
        self.assert_scenario('unrelated-attachment', False)
    def test_sent_message_with_linked_valid_attachment_passes(self):
        self.assert_scenario('valid', True)
    def test_pg_backend_generic_mime_is_supported(self):
        self.assert_scenario('generic-mime', True)
    def test_other_checks_remain_partial_instead_of_all_becoming_mandatory(self):
        self.assert_scenario('partial-body', True)
    def test_mandatory_evidence_does_not_override_low_numeric_score(self):
        for spec in SPECS:
            with self.subTest(task=spec[0]):
                code, result, text = self.evaluate(spec, 'valid', low_score=True)
                self.assertEqual(code, 1, text)
                self.assertTrue(result['mandatory_pass'])
                self.assertLess(result['accuracy'], 70)
                self.assertFalse(result['pass'])


if __name__ == '__main__':
    unittest.main()
