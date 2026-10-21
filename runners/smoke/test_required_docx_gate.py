"""Smoke tests for the fail-closed required-DOCX gate.

Covers utils/required_artifacts.require_readable_docx directly and the six
finalpool graders that must invoke it right after parse_args(), before any
DB reads/writes or component checks:

    python3 -m unittest discover -s runners/smoke -p test_required_docx_gate.py -v
"""
import importlib.util
import io
import sys
import zipfile
from contextlib import redirect_stdout
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase, main
from unittest.mock import patch

REPO = Path(__file__).resolve().parents[2]
if str(REPO) not in sys.path:
    sys.path.insert(0, str(REPO))

from docx import Document

from utils.required_artifacts import require_readable_docx

TASKS = [
    ("arxiv-latex-analysis-teamly-word", "Paper_Analysis_Report.docx"),
    ("insales-coupon-analysis", "Coupon_Strategy.docx"),
    ("rzd-kulinar-novgorod-trip-gcal-word", "Novgorod_Trip_Itinerary.docx"),
    ("rzd-msk-spb-trip-notion-gcal-word", "Travel_Plan.docx"),
    ("rzd-team-retreat-teamly-word-gcal", "Team_Retreat_Itinerary.docx"),
    ("terminal-sf-insales-order-audit-excel-word", "Audit_Findings.docx"),
]


def load_grader(task_dir):
    path = REPO / "tasks" / "finalpool" / task_dir / "evaluation" / "main.py"
    spec = importlib.util.spec_from_file_location(f"gate_{task_dir}", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


class TestHelper(TestCase):
    def test_accepts_genuine_docx_and_preserves_bytes(self):
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "genuine.docx"
            doc = Document()
            doc.add_paragraph("genuine content")
            doc.save(path)
            before = path.read_bytes()
            out = io.StringIO()
            with redirect_stdout(out):
                parsed = require_readable_docx(path)
            self.assertTrue(any(p.text == "genuine content" for p in parsed.paragraphs))
            self.assertEqual(path.read_bytes(), before)

    def test_rejects_missing_file(self):
        with TemporaryDirectory() as tmp:
            out = io.StringIO()
            with redirect_stdout(out), self.assertRaises(SystemExit) as cm:
                require_readable_docx(Path(tmp) / "absent.docx")
            self.assertEqual(cm.exception.code, 1)
            self.assertIn("CRITICAL", out.getvalue())

    def _reject_bytes(self, name, payload):
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / name
            path.write_bytes(payload)
            out = io.StringIO()
            with redirect_stdout(out), self.assertRaises(SystemExit) as cm:
                require_readable_docx(path)
            self.assertEqual(cm.exception.code, 1)
            self.assertIn("CRITICAL", out.getvalue())
            self.assertEqual(path.read_bytes(), payload)  # untouched

    def test_rejects_plain_text_renamed_docx(self):
        self._reject_bytes("plain.docx", "не ZIP, просто текст\n".encode("utf-8"))

    def test_rejects_malformed_zip(self):
        self._reject_bytes("broken.docx", b"PK\x03\x04 truncated garbage")

    def test_rejects_valid_zip_without_docx_payload(self):
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "notdocx.docx"
            with zipfile.ZipFile(path, "w") as zf:
                zf.writestr("readme.txt", "no OOXML parts here")
            out = io.StringIO()
            with redirect_stdout(out), self.assertRaises(SystemExit) as cm:
                require_readable_docx(path)
            self.assertEqual(cm.exception.code, 1)
            self.assertIn("CRITICAL", out.getvalue())


class TestGraders(TestCase):
    """Each real main() must exit 1 at the gate, before any DB use."""

    def _run_grader(self, task_dir, filename, fake_bytes=None):
        mod = load_grader(task_dir)
        with TemporaryDirectory() as tmp:
            agent = Path(tmp) / "agent"
            agent.mkdir()
            gt = Path(tmp) / "gt"
            gt.mkdir()
            target = agent / filename
            if fake_bytes is not None:
                target.write_bytes(fake_bytes)
            before = target.read_bytes() if target.exists() else None
            out = io.StringIO()
            argv = ["main.py", "--agent_workspace", str(agent),
                    "--groundtruth_workspace", str(gt)]
            with patch.object(sys, "argv", argv), \
                 patch("psycopg2.connect",
                       side_effect=AssertionError("DB used before required-docx gate")), \
                 redirect_stdout(out):
                with self.assertRaises(SystemExit) as cm:
                    mod.main()
            self.assertEqual(cm.exception.code, 1, out.getvalue())
            self.assertIn("CRITICAL", out.getvalue())      # critical diagnostic
            self.assertNotIn("PASS", out.getvalue())       # no quality PASS
            after = target.read_bytes() if target.exists() else None
            self.assertEqual(before, after)                # no overwrite
            expected = [target] if target.exists() else []
            self.assertEqual(sorted(agent.iterdir()), sorted(expected))  # no new files

    def test_graders_reject_plain_text_docx(self):
        for task_dir, filename in TASKS:
            with self.subTest(task=task_dir):
                self._run_grader(task_dir, filename, b"plain text, not a zip\n")

    def test_graders_reject_missing_docx(self):
        for task_dir, filename in TASKS:
            with self.subTest(task=task_dir):
                self._run_grader(task_dir, filename)


if __name__ == "__main__":
    main()
