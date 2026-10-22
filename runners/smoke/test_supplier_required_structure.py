"""Actual supplier Excel checker/main; calendar and email counts are isolated controls."""
from pathlib import Path
import contextlib
import importlib.util
import io
import json
import sys
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[2]
TASK = 'rzd-insales-elektrosnab-supplier-visit-spb-kzn-excel-email-gcal'


class SupplierGateTests(unittest.TestCase):
    def fixture(self, folder, wrong_priority=False, wrong_prices=False,
                missing_headers=None, absent_sheet=None):
        book = Workbook()
        products = book.active
        products.title = 'Products'
        headers = ['Product_ID', 'Name', 'SKU', 'Current_Stock', 'Monthly_Sales',
                   'Supplier_Name', 'Supplier_City', 'Priority']
        products.append(['col'+str(i) for i in range(8)] if missing_headers == 'Products' else headers)
        for i in range(5):
            products.append([str(i), 'item', 'sku', 1, 10, 'supplier'+str(i),
                             'Санкт-Петербург' if i < 3 else 'Казань',
                             'Low' if wrong_priority else 'High'])
        travel = book.create_sheet('Travel_Plan')
        travel.append(['a', 'b', 'c', 'd', 'e', 'f', 'g'] if missing_headers == 'Travel_Plan'
                      else ['Supplier_City', 'Train_No', 'Departure_Time', 'Arrival_Time',
                            'Duration_Min', 'Seat_Class', 'Ticket_Price_RUB'])
        travel.append(['Санкт-Петербург', '752А', '06:50', '10:50', 240, 'Бизнес',
                       1 if wrong_prices else 14000])
        travel.append(['Казань', '716Г', '06:25', '11:50', 325, 'Бизнес',
                       1 if wrong_prices else 11000])
        visits = book.create_sheet('Visit_Schedule')
        visits.append(['Date', 'Supplier_Name', 'City', 'Meeting_Time', 'Purpose', 'Status'])
        for i in range(5):
            visits.append(['2026-03-10' if i < 3 else '2026-03-17', 'supplier'+str(i),
                           'Санкт-Петербург' if i < 3 else 'Казань', '13:00', 'meeting', 'Scheduled'])
        if absent_sheet:
            del book[absent_sheet]
        book.save(folder / 'Supplier_Visit_Plan.xlsx')
        book.close()

    def evaluate(self, folder):
        path = ROOT / 'tasks/finalpool' / TASK / 'evaluation/main.py'
        spec = importlib.util.spec_from_file_location('supplier_gate_control', path)
        grader = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(grader)
        def positive_controls(n):
            for i in range(n):
                grader.record('synthetic service check '+str(i), True)
        # These counts correspond to the three GCal and four email checks.
        # We deliberately isolate the real Excel checker and final decision.
        grader.check_gcal = lambda: positive_controls(3)
        grader.check_emails = lambda: positive_controls(4)
        result_file = folder / 'result.json'
        with patch.object(sys, 'argv', ['grader', '--agent_workspace', str(folder),
                                      '--res_log_file', str(result_file)]), \
             patch.object(grader.psycopg2, 'connect', side_effect=AssertionError('DB forbidden')), \
             contextlib.redirect_stdout(io.StringIO()):
            with self.assertRaises(SystemExit) as exited:
                grader.main()
        return exited.exception.code, json.loads(result_file.read_text())

    def test_valid_deliverable_passes(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            self.fixture(folder)
            code, result = self.evaluate(folder)
            self.assertEqual(code, 0)
            self.assertEqual(result['accuracy'], 100)
            self.assertFalse(result['critical_fails'])

    def test_wrong_priority_fails_with_or_without_headers(self):
        for missing in [None, 'Products']:
            with self.subTest(missing=missing), tempfile.TemporaryDirectory() as tmp:
                folder = Path(tmp)
                self.fixture(folder, wrong_priority=True, missing_headers=missing)
                code, result = self.evaluate(folder)
                self.assertGreaterEqual(result['accuracy'], 70)
                self.assertEqual(code, 1)
                self.assertTrue(result['critical_fails'])

    def test_wrong_prices_fail_with_or_without_headers(self):
        for missing in [None, 'Travel_Plan']:
            with self.subTest(missing=missing), tempfile.TemporaryDirectory() as tmp:
                folder = Path(tmp)
                self.fixture(folder, wrong_prices=True, missing_headers=missing)
                code, result = self.evaluate(folder)
                self.assertGreaterEqual(result['accuracy'], 70)
                self.assertEqual(code, 1)
                self.assertTrue(result['critical_fails'])

    def test_required_sheets_cannot_disable_critical_checks(self):
        for sheet in ['Products', 'Travel_Plan']:
            with self.subTest(sheet=sheet), tempfile.TemporaryDirectory() as tmp:
                folder = Path(tmp)
                self.fixture(folder, absent_sheet=sheet)
                code, result = self.evaluate(folder)
                self.assertGreaterEqual(result['accuracy'], 70)
                self.assertEqual(code, 1)
                self.assertTrue(result['critical_fails'])

    def test_empty_required_sheets_cannot_skip_critical_checks(self):
        for sheet in ['Products', 'Travel_Plan']:
            with self.subTest(sheet=sheet), tempfile.TemporaryDirectory() as tmp:
                folder = Path(tmp)
                self.fixture(folder)
                path = folder / 'Supplier_Visit_Plan.xlsx'
                book = load_workbook(path)
                book[sheet].delete_rows(1, book[sheet].max_row)
                book.save(path)
                book.close()
                code, result = self.evaluate(folder)
                self.assertEqual(code, 1)
                self.assertTrue(result['critical_fails'])

    def test_missing_workbook_cannot_be_offset_by_service_points(self):
        with tempfile.TemporaryDirectory() as tmp:
            code, result = self.evaluate(Path(tmp))
            self.assertGreaterEqual(result['accuracy'], 70)
            self.assertEqual(code, 1)
            self.assertTrue(result['critical_fails'])

    def test_unreadable_workbook_cannot_be_offset_by_service_points(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            (folder / 'Supplier_Visit_Plan.xlsx').write_bytes(b'not an Office package')
            code, result = self.evaluate(folder)
            self.assertGreaterEqual(result['accuracy'], 70)
            self.assertEqual(code, 1)
            self.assertTrue(result['critical_fails'])

    def test_good_values_still_need_enabling_headers(self):
        for sheet in ['Products', 'Travel_Plan']:
            with self.subTest(sheet=sheet), tempfile.TemporaryDirectory() as tmp:
                folder = Path(tmp)
                self.fixture(folder, missing_headers=sheet)
                code, result = self.evaluate(folder)
                self.assertGreaterEqual(result['accuracy'], 70)
                self.assertEqual(code, 1)
                self.assertTrue(result['critical_fails'])


if __name__ == '__main__':
    unittest.main()
