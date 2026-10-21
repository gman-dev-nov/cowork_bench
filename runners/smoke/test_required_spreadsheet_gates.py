"""Fail-closed mandatory spreadsheet structure, with real Office parsers."""

import contextlib
import importlib.util
import io
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook
from utils.required_artifacts import require_readable_xlsx

ROOT = Path(__file__).resolve().parents[2]


def module(task):
    spec = importlib.util.spec_from_file_location(
        task.replace("-", "_"), ROOT / "tasks/finalpool" / task / "evaluation/main.py"
    )
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


class SpreadsheetGates(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_required_xlsx_gate_before_db_for_absent_plain_text_and_wrong_zip(self):
        m = module("sf-support-resolution-gsheet")
        p = self.root / "Support_Resolution_Analysis.xlsx"
        for content in [None, b"CSV,not,Office", b"PK malformed"]:
            if content is None:
                p.unlink(missing_ok=True)
            else:
                p.write_bytes(content)
            with (
                self.subTest(content=content),
                patch.object(
                    sys, "argv", ["grader", "--agent_workspace", str(self.root)]
                ),
                patch.object(m.psycopg2, "connect") as db,
                contextlib.redirect_stdout(io.StringIO()),
                self.assertRaises(SystemExit) as exc,
            ):
                m.main()
            self.assertEqual(exc.exception.code, 1)
            db.assert_not_called()
        book = Workbook()
        book.save(p)
        before = p.read_bytes()
        require_readable_xlsx(p)
        self.assertEqual(p.read_bytes(), before)

    def test_inventory_missing_headers_report_critical_failure_without_crashing(self):
        import datetime

        for mutation in ("stock_headers", "schedule_headers", "unreadable"):
            with (
                self.subTest(mutation=mutation),
                contextlib.redirect_stdout(io.StringIO()),
            ):
                m = module("insales-inventory-reorder-forecast")
                book = Workbook()
                sheet = book.active
                sheet.title = "Stock Analysis"
                headers = [
                    "Product_ID",
                    "Product_Name",
                    "Category",
                    "Current_Stock",
                    "Avg_Daily_Sales",
                    "Days_Until_Safety_Stock",
                    "Lead_Time",
                    "Needs_Reorder",
                    "Reorder_Date",
                ]
                sheet.append(["wrong"] if mutation == "stock_headers" else headers)
                sheet.append([1, "Product", "Аудио", 10, 1, 5, 14, "Yes", "2030-01-02"])
                schedule = book.create_sheet("Reorder Schedule")
                schedule.append(
                    ["wrong"]
                    if mutation == "schedule_headers"
                    else [
                        "Product_Name",
                        "Category",
                        "Current_Stock",
                        "Reorder_Date",
                        "Lead_Time_Days",
                    ]
                )
                schedule.append(["Product", "Аудио", 10, "2030-01-02", 14])
                path = self.root / "Inventory_Forecast.xlsx"
                book.save(path)
                if mutation == "unreadable":
                    path.write_text("not a workbook")
                with patch.object(m, "get_gcal_events", return_value=[]):
                    m.evaluate(
                        str(self.root),
                        {"total_products": 1, "reorder_count": 1, "safety_stock": 5},
                        datetime.date(2030, 1, 1),
                    )
                self.assertTrue(set(m.FAILED_NAMES) & m.CRITICAL_CHECKS)

    def make_finance(self, m):
        rows = []
        for i, sym in enumerate(["A", "B"]):
            rows.append(
                dict(
                    Symbol=sym,
                    Company_Name=sym,
                    Total_Revenue=100,
                    Net_Income=20,
                    Gross_Profit=30,
                    Operating_Income=10,
                    Profit_Margin_Pct=20,
                    Operating_Margin_Pct=10,
                    Credit_Rating="AA",
                    Rating_Outlook="Stable",
                    Margin_Health="Strong",
                )
            )
        book = Workbook()
        book.remove(book.active)
        for name, cols in [
            (
                "Income Analysis",
                [
                    "Symbol",
                    "Company_Name",
                    "Total_Revenue",
                    "Net_Income",
                    "Gross_Profit",
                    "Operating_Income",
                    "Profit_Margin_Pct",
                    "Operating_Margin_Pct",
                ],
            ),
            (
                "Health Assessment",
                [
                    "Symbol",
                    "Credit_Rating",
                    "Rating_Outlook",
                    "Profit_Margin_Pct",
                    "Margin_Health",
                ],
            ),
        ]:
            sheet = book.create_sheet(name)
            sheet.append(cols)
            for row in rows:
                sheet.append([row[c] for c in cols])
        book.create_sheet("Summary").append(["Metric", "Value"])
        return book, rows

    def test_finance_valid_structure_and_missing_header_row_company_sheet_or_file(self):
        for mutation in [
            "valid",
            "no_header",
            "no_row",
            "duplicate_company",
            "no_income",
            "no_health",
            "no_summary",
            "absent",
            "bad_zip",
        ]:
            with (
                self.subTest(mutation=mutation),
                contextlib.redirect_stdout(io.StringIO()),
            ):
                m = module("moex-financial-health-gsheet")
                book, rows = self.make_finance(m)
                if mutation == "no_header":
                    book["Income Analysis"].delete_rows(1)
                if mutation == "no_row":
                    book["Income Analysis"].delete_rows(3)
                if mutation == "duplicate_company":
                    book["Health Assessment"]["A3"] = "A"
                for key, title in [
                    ("no_income", "Income Analysis"),
                    ("no_health", "Health Assessment"),
                    ("no_summary", "Summary"),
                ]:
                    if mutation == key:
                        del book[title]
                path = self.root / "financial_health_report.xlsx"
                book.save(path)
                if mutation == "absent":
                    path.unlink()
                if mutation == "bad_zip":
                    path.write_text("not xlsx")
                passed, total = m.check_local_excel(self.root, rows, {})
                if mutation == "valid":
                    self.assertFalse(m.CRITICAL_FAILURES)
                    self.assertEqual(passed, total)
                else:
                    self.assertTrue(m.CRITICAL_FAILURES, mutation)

    def test_arxiv_required_headers_do_not_disable_title_checks(self):
        for mutation in [
            "valid",
            "no_title",
            "no_summary_header",
            "missing_sheet",
            "unreadable",
        ]:
            with (
                self.subTest(mutation=mutation),
                contextlib.redirect_stdout(io.StringIO()),
            ):
                m = module("arxiv-federated-learning-ru-teamly-excel")
                book = Workbook()
                sheet = book.active
                sheet.title = "Paper Details"
                sheet.append(
                    ["Paper_ID", "Title", "Authors", "Year", "Abstract_Length"]
                )
                for pid, title in m.EXPECTED_FL_PAPERS.items():
                    sheet.append([pid, title, "Author", 2020, 100])
                summary = book.create_sheet("Summary")
                summary.append(["Metric", "Value"])
                for row in [
                    ("Total_Papers", 4),
                    ("Avg_Abstract_Length", 100),
                    ("Earliest_Year", 2020),
                    ("Latest_Year", 2020),
                ]:
                    summary.append(row)
                if mutation == "no_title":
                    sheet["B1"] = "unrelated"
                if mutation == "no_summary_header":
                    summary.delete_rows(1)
                if mutation == "missing_sheet":
                    del book["Paper Details"]
                path = self.root / "Federated_Learning_Papers.xlsx"
                book.save(path)
                if mutation == "unreadable":
                    path.write_text("CSV,not,Office")
                m.check_excel(self.root)
                critical = set(m.FAILED_NAMES) & m.CRITICAL_CHECKS
                if mutation == "valid":
                    self.assertFalse(critical)
                else:
                    self.assertTrue(critical, mutation)


if __name__ == "__main__":
    unittest.main()
