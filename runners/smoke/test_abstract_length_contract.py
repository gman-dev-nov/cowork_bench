"""Real grader main and XLSX; narrow read-only DB seam, no live connection."""
from pathlib import Path
import contextlib
import importlib.util
import io
import json
import sys
import tempfile
import unittest
from unittest.mock import patch
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
TASK = ROOT / 'tasks/finalpool/arxiv-federated-learning-ru-teamly-excel'


def load(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class AbstractLengthContractTests(unittest.TestCase):
    def evaluate(self, change=None, canonical=None, source_error=None, invalid_file=None):
        seed = load(TASK / 'preprocess/main.py', 'abstract_contract_seed').FL_PAPERS
        source = canonical if canonical is not None else {p['id']: p['summary'] for p in seed}
        # Build counts from the known public seed; source can independently fail.
        texts = source if canonical is not None and all(isinstance(x, str) for x in source.values()) else {p['id']: p['summary'] for p in seed}
        if len(texts) != 4:
            texts = {p['id']: p['summary'] for p in seed}
        lengths = {p['id']: len(texts[p['id']]) for p in seed}
        mean = round(sum(lengths.values()) / 4)
        grader = load(TASK / 'evaluation/main.py', 'abstract_contract_grader')
        class Cursor:
            def execute(self, sql):
                if 'FROM teamly.pages' in sql:
                    self.rows = [(4,'База знаний по федеративному обучению','hub',None)] + [
                        (i,p['title'],'federated learning',4) for i,p in enumerate(seed,5)]
                elif 'FROM arxiv.papers' in sql:
                    if source_error:
                        raise RuntimeError(source_error)
                    self.rows = list(source.items())
                else:
                    raise AssertionError('Unexpected SQL: '+sql)
            def fetchall(self): return self.rows
            def close(self): pass
        class Connection:
            def cursor(self): return Cursor()
            def close(self): pass
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp); path = folder / 'Federated_Learning_Papers.xlsx'
            book = openpyxl.Workbook(); details = book.active; details.title = 'Paper Details'
            details.append(['Paper_ID','Title','Authors','Year','Abstract_Length'])
            for p in seed:
                details.append([p['id'],p['title'],'Synthetic author',int(p['published'][:4]),lengths[p['id']]])
            summary = book.create_sheet('Summary'); summary.append(['Metric','Value'])
            for k,v in [('Total_Papers',4),('Avg_Abstract_Length',mean),('Earliest_Year',2016),('Latest_Year',2019)]:
                summary.append([k,v])
            if change: change(book)
            book.save(path); book.close()
            if invalid_file == 'missing': path.unlink()
            elif invalid_file == 'corrupt': path.write_bytes(b'not an XLSX package')
            before = path.read_bytes() if path.exists() else None
            result_path = folder / 'result.json'
            with patch.object(grader.psycopg2,'connect',side_effect=lambda **kw: Connection()), \
                 patch.object(sys,'argv',['grader','--agent_workspace',str(folder),'--res_log_file',str(result_path)]), \
                 contextlib.redirect_stdout(io.StringIO()):
                with self.assertRaises(SystemExit) as exited: grader.main()
            self.assertEqual(path.read_bytes() if path.exists() else None, before)
            return exited.exception.code,json.loads(result_path.read_text())

    def rejected(self, **kwargs):
        code,result = self.evaluate(**kwargs)
        self.assertEqual(code,1,result)
        self.assertTrue(result['critical_failed'],result)
        return result

    def test_valid_numeric_text_and_integral_float(self):
        for convert in [lambda x:x,str,float]:
            def change(book):
                for row in list(book['Paper Details'].rows)[1:]: row[4].value=convert(row[4].value)
                book['Summary'].cell(3,2).value=convert(book['Summary'].cell(3,2).value)
            code,result = self.evaluate(change=change)
            self.assertEqual(code,0,result); self.assertFalse(result['critical_failed'])

    def test_invalid_lengths_cannot_be_offset_by_other_points(self):
        for value in [None,-999,True,2.5,'NaN','inf','bad',1231]:
            with self.subTest(value=value):
                result=self.rejected(change=lambda b: setattr(b['Paper Details'].cell(2,5),'value',value))
                self.assertGreaterEqual(result['accuracy'],80)

    def test_invalid_means_are_diagnostic_failures_with_result_json(self):
        for value in [None,-1,True,804.5,'NaN','Infinity','bad',805]:
            with self.subTest(value=value):
                result=self.rejected(change=lambda b: setattr(b['Summary'].cell(3,2),'value',value))
                self.assertGreaterEqual(result['accuracy'],80)

    def test_missing_duplicate_unknown_and_structural_controls(self):
        def duplicate(book): book['Paper Details'].cell(5,1).value=book['Paper Details'].cell(2,1).value
        def unknown(book): book['Paper Details'].cell(5,1).value='9999.99999'
        def missing_row(book): book['Paper Details'].delete_rows(5)
        def missing_id(book): book['Paper Details'].cell(1,1).value='Broken'
        def missing_length(book): book['Paper Details'].cell(1,5).value='Broken'
        def missing_sheet(book): del book['Summary']
        def missing_mean(book): book['Summary'].delete_rows(3)
        def duplicate_mean(book): book['Summary'].append(['Avg_Abstract_Length',804])
        for change in [duplicate,unknown,missing_row,missing_id,missing_length,missing_sheet,missing_mean,duplicate_mean]:
            with self.subTest(change=change.__name__): self.rejected(change=change)

    def test_canonical_source_is_required_and_failure_has_no_agent_fallback(self):
        seed=load(TASK/'preprocess/main.py','abstract_canonical_seed').FL_PAPERS
        canonical={p['id']:p['summary'] for p in seed}
        self.rejected(source_error='SYNTHETIC_DB_FAILURE')
        self.rejected(canonical={k:v for k,v in canonical.items() if k!='1602.05629'})
        self.rejected(canonical={**canonical,'1602.05629':None})

    def test_unicode_codepoints_and_explicit_ties_to_even(self):
        seed=load(TASK/'preprocess/main.py','abstract_unicode_seed').FL_PAPERS
        texts={p['id']:('Ж😀e\u0301\n'*50)[:201+i] for i,p in enumerate(seed)}
        code,result=self.evaluate(canonical=texts)
        self.assertEqual(code,0,result)  # round(202.5) == 202
        def bytes_instead_of_characters(book):
            for row in list(book['Paper Details'].rows)[1:]: row[4].value=len(texts[row[0].value].encode('utf-8'))
        self.rejected(canonical=texts,change=bytes_instead_of_characters)
        self.rejected(canonical=texts,change=lambda b:setattr(b['Summary'].cell(3,2),'value',203))

    def test_missing_and_corrupt_workbook_fail_without_rewriting(self):
        for kind in ['missing','corrupt']:
            with self.subTest(kind=kind): self.rejected(invalid_file=kind)


if __name__=='__main__': unittest.main()
